#!/usr/bin/env python3
# This file is part of FDTP (File & Directory Tool for Pomera).
# SPDX-FileCopyrightText: 2026 yasoteien
# SPDX-License-Identifier: AGPL-3.0-or-later
"""Makes the user manual's Word files from the Markdown sources.

    python3 fdtp/tools/make_manual.py [ja|en ...]

Reads docs/manual_src/FDTP_manual_{ja,en}.md (images relative to that
directory) and the version from fdtp/src/version.h, and writes
fdtp/public/manual/{ja,en}/FDTP_manual_{ja,en}.docx (make_public.sh copies
them to docs/manual/ of the public tree). The PDF is made from the .docx in
Microsoft Word.

Needs python-docx (python3 -m pip install --user python-docx).

The Markdown understood is the subset the sources use: "## N." chapters
(Heading 1), "###" (Heading 2), "####" (Heading 3), paragraphs, "- " lists
with "  1." items under them, "1." lists, pipe tables (cells may hold an
image), fenced code, a line holding only an image (caption = alt text),
`code`, **bold** and bare https:// URLs (made into hyperlinks). Everything
before the first "## " (the title lines) is replaced by the cover; "---"
lines are dropped. The cover photo is not repeated in the body.
"""
import datetime
import os
import re
import struct
import sys

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK, WD_LINE_SPACING
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)                      # fdtp/
ROOT = os.path.dirname(APP)                      # the repository
SRC = os.path.join(ROOT, "docs", "manual_src")
OUT = os.path.join(APP, "public", "manual")
COVER_PHOTO = "pomera-filer.jpg"     # not repeated in the body
COVER_ART = "cover-art.jpg"          # docs/media/fdtp/manual/ (generated illustration)

PAGE_W, MARGIN = 210, 22                         # mm, A4 portrait
TEXT_W = PAGE_W - 2 * MARGIN                     # 166 mm
SHOT_W, PHOTO_W, MAX_H = 150, 120, 110           # image widths / height cap (mm)

LANG = {
    "ja": dict(
        title="FDTP 利用の手引き",
        sub="File & Directory Tool for Pomera",
        ver="版 {v}", author="作者 yasoteien",
        blurb=["✨ ポメラ純正ソフトにないから、ルビ振り対応の縦横表示エディタ作りました！ ✨",
               "（なろう・カクヨム対応、出版社向け整形付き）",
               "あ、ついでに画像・Excel・Word・PowerPoint・PDF も見えるファイラーもつけたよ！",
               "ついでのついでに、PDF の結合・分割・書き込みもつけたよ！",
               "パワポは F5 でスライドショーもできるよ！",
               "英語の画面もつけたよ！（US 版ポメラの人へ）🎉"],
        toc="目次", toc_hint="（目次は、Word で開いた時に更新されます。更新されない時は、ここを右クリックして「フィールド更新」を選んでください）",
        fig="図 {n}　{t}",
        body="游明朝", head="游ゴシック", latin_body="游明朝", latin_head="游ゴシック",
        mono="Consolas", mono_ea="游ゴシック",
        size=10.5, exact=17.5, cell=9.5, cell_exact=14, lang="ja-JP"),
    "en": dict(
        title="FDTP User Manual",
        sub="File & Directory Tool for Pomera",
        ver="Version {v}", author="By yasoteien",
        blurb=["✨ The stock Pomera software didn't have one, so I made a vertical & horizontal novel editor with ruby! ✨",
               "(Narou and Kakuyomu notation, publisher-style formatting included)",
               "Oh, and while I was at it: a file manager that shows images, Excel, Word, PowerPoint and PDF too!",
               "And while I was at that: merge, split and mark up PDFs too!",
               "PowerPoint? Press F5 for a slide show!",
               "And English screens, for Pomera DM250US folks! 🎉"],
        toc="Contents", toc_hint="(The table of contents is updated when Word opens the file. If it is not, right-click here and choose Update Field.)",
        fig="Figure {n}. {t}",
        body="Calibri", head="Calibri", latin_body="Calibri", latin_head="Calibri",
        mono="Consolas", mono_ea="Consolas",
        size=10.5, exact=None, cell=9.5, cell_exact=None, lang="en-US"),
}

HEAD_COLOR = RGBColor(0x1F, 0x3A, 0x5F)
LINK_COLOR = RGBColor(0x05, 0x63, 0xC1)
HEADER_FILL = "DCE6F2"
CODE_FILL = "F2F2F2"


def version():
    s = open(os.path.join(APP, "src", "version.h"), encoding="utf-8").read()
    m = re.search(r'#define\s+FDTP_VERSION\s+"([^"]+)"', s)
    if not m:
        sys.exit("make_manual: no FDTP_VERSION in src/version.h")
    return m.group(1)


# ---- small XML helpers ---------------------------------------------------------

def set_fonts(rpr_parent, ascii_, east):
    """Sets w:rFonts on an rPr (dropping the theme fonts, which win otherwise)."""
    rpr = rpr_parent.get_or_add_rPr() if hasattr(rpr_parent, "get_or_add_rPr") else rpr_parent
    old = rpr.find(qn("w:rFonts"))
    if old is not None:
        rpr.remove(old)
    f = OxmlElement("w:rFonts")
    for k, v in (("w:ascii", ascii_), ("w:hAnsi", ascii_), ("w:eastAsia", east), ("w:cs", ascii_)):
        f.set(qn(k), v)
    rpr.insert(0, f)


def shade(el_pr, fill):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")                # ShadingType.CLEAR
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    el_pr.append(shd)


def image_px(path):
    """(width, height) in pixels of a PNG or JPEG."""
    with open(path, "rb") as f:
        head = f.read(26)
        if head[:8] == b"\x89PNG\r\n\x1a\n":
            return struct.unpack(">II", head[16:24])
        f.seek(2)
        while True:
            b = f.read(1)
            while b and b != b"\xff":
                b = f.read(1)
            while b == b"\xff":
                b = f.read(1)
            if not b:
                break
            marker = b[0]
            ln = struct.unpack(">H", f.read(2))[0]
            if marker in (0xC0, 0xC1, 0xC2):
                f.read(1)
                h, w = struct.unpack(">HH", f.read(4))
                return w, h
            f.seek(ln - 2, 1)
    sys.exit(f"make_manual: cannot read the size of {path}")


def image_size(path, want_w):
    w, h = image_px(path)
    width = want_w
    if width * h / w > MAX_H:
        width = MAX_H * w / h
    return Mm(width)


# ---- the document ---------------------------------------------------------------

class Builder:
    def __init__(self, lang, ver):
        self.L = LANG[lang]
        self.lang = lang
        self.ver = ver
        self.fig = 0
        self.doc = Document()
        self.setup()

    # styles, page, settings
    def setup(self):
        L, d = self.L, self.doc
        sec = d.sections[0]
        sec.orientation = WD_ORIENT.PORTRAIT
        sec.page_width, sec.page_height = Mm(210), Mm(297)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Mm(MARGIN))
        sec.top_margin, sec.bottom_margin = Mm(25), Mm(22)
        sec.header_distance, sec.footer_distance = Mm(12), Mm(10)
        sec.different_first_page_header_footer = True
        # page number in the footer (not on the cover)
        p = sec.footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        self.field(p, "PAGE", "1")

        # document defaults: fonts and language
        styles = d.styles.element
        rpr_default = styles.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
        set_fonts(rpr_default, L["latin_body"], L["body"])
        lang_el = rpr_default.find(qn("w:lang"))
        if lang_el is None:
            lang_el = OxmlElement("w:lang")
            rpr_default.append(lang_el)
        lang_el.set(qn("w:val"), "en-US")
        lang_el.set(qn("w:eastAsia"), "ja-JP")
        # Word 2013 and later (python-docx's template opens in compatibility mode)
        st = d.settings.element
        compat = st.find(qn("w:compat"))
        if compat is None:
            compat = OxmlElement("w:compat")
            st.append(compat)
        for cs in compat.findall(qn("w:compatSetting")):
            if cs.get(qn("w:name")) == "compatibilityMode":
                compat.remove(cs)
        cs = OxmlElement("w:compatSetting")
        cs.set(qn("w:name"), "compatibilityMode")
        cs.set(qn("w:uri"), "http://schemas.microsoft.com/office/word")
        cs.set(qn("w:val"), "15")
        compat.append(cs)

        normal = d.styles["Normal"]
        normal.font.size = Pt(L["size"])
        set_fonts(normal.element.get_or_add_rPr(), L["latin_body"], L["body"])
        pf = normal.paragraph_format
        pf.space_before, pf.space_after = Pt(0), Pt(6)
        if L["exact"]:
            pf.line_spacing_rule = WD_LINE_SPACING.EXACTLY
            pf.line_spacing = Pt(L["exact"])
        else:
            pf.line_spacing = 1.2

        sizes = {1: 16, 2: 13, 3: 11.5}
        for lvl, sz in sizes.items():
            st = d.styles[f"Heading {lvl}"]
            st.font.size = Pt(sz)
            st.font.bold = True
            st.font.italic = False
            st.font.color.rgb = HEAD_COLOR
            set_fonts(st.element.get_or_add_rPr(), L["latin_head"], L["head"])
            f = st.paragraph_format
            f.space_before = Pt({1: 0, 2: 14, 3: 10}[lvl])
            f.space_after = Pt({1: 12, 2: 6, 3: 4}[lvl])
            f.keep_with_next = True
            f.line_spacing_rule = WD_LINE_SPACING.SINGLE
            f.page_break_before = lvl == 1
        cap = d.styles["Caption"]
        cap.font.size = Pt(9)
        cap.font.bold = False
        cap.font.italic = False
        cap.font.color.rgb = RGBColor(0x40, 0x40, 0x40)
        set_fonts(cap.element.get_or_add_rPr(), L["latin_head"], L["head"])
        cap.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
        cap.paragraph_format.space_after = Pt(10)

        # update the fields (the TOC) when Word opens the file
        settings = d.settings.element
        uf = OxmlElement("w:updateFields")
        uf.set(qn("w:val"), "true")
        settings.append(uf)

        cp = d.core_properties
        cp.title = L["title"]
        cp.author = "yasoteien"
        cp.last_modified_by = "yasoteien"
        cp.language = L["lang"]
        cp.subject = ""
        cp.keywords = "FDTP"
        cp.comments = ""
        cp.created = cp.modified = datetime.datetime.now(datetime.timezone.utc).replace(microsecond=0)
        # the default template's thumbnail is not this document's
        pkg = d.part.package
        for rid, rel in list(pkg.rels.items()):
            if rel.reltype.endswith("/thumbnail"):
                del pkg.rels[rid]

    def field(self, p, instr, placeholder):
        """A complex field (begin, instruction, separate, result, end)."""
        def fc(kind):
            r = p.add_run()
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            if kind == "begin":
                el.set(qn("w:dirty"), "true")
            r._r.append(el)
        fc("begin")
        r = p.add_run()
        t = OxmlElement("w:instrText")
        t.set(qn("xml:space"), "preserve")
        t.text = f" {instr} "
        r._r.append(t)
        fc("separate")
        p.add_run(placeholder)
        fc("end")

    # inline text: `code`, **bold**, URLs
    INLINE = re.compile(r"`([^`]+)`|\*\*(.+?)\*\*|(https?://[A-Za-z0-9/._:%?=&#~+\-]+)")

    def inline(self, p, text, size=None, bold=False):
        pos = 0
        for m in self.INLINE.finditer(text):
            if m.start() > pos:
                self.run(p, text[pos:m.start()], size, bold)
            if m.group(1) is not None:
                self.run(p, m.group(1), size, bold, mono=True)
            elif m.group(2) is not None:
                self.inline(p, m.group(2), size, True)
            else:
                url = m.group(3)
                tail = ""
                while url and url[-1] in ".,":
                    tail = url[-1] + tail
                    url = url[:-1]
                self.link(p, url, url, size)
                if tail:
                    self.run(p, tail, size, bold)
            pos = m.end()
        if pos < len(text):
            self.run(p, text[pos:], size, bold)

    def run(self, p, text, size=None, bold=False, mono=False):
        r = p.add_run(text)
        if size:
            r.font.size = Pt(size)
        if bold:
            r.font.bold = True
        if mono:
            set_fonts(r._r.get_or_add_rPr(), self.L["mono"], self.L["mono_ea"])
            r.font.size = Pt((size or self.L["size"]) - 0.5)
        return r

    def link(self, p, url, text, size=None):
        rid = p.part.relate_to(url, RT.HYPERLINK, is_external=True)
        h = OxmlElement("w:hyperlink")
        h.set(qn("r:id"), rid)
        h.set(qn("w:history"), "1")
        r = OxmlElement("w:r")
        rpr = OxmlElement("w:rPr")
        c = OxmlElement("w:color")
        c.set(qn("w:val"), str(LINK_COLOR))
        u = OxmlElement("w:u")
        u.set(qn("w:val"), "single")
        rpr.append(c)
        rpr.append(u)
        if size:
            for tag in ("w:sz", "w:szCs"):
                s = OxmlElement(tag)
                s.set(qn("w:val"), str(int(size * 2)))
                rpr.append(s)
        r.append(rpr)
        t = OxmlElement("w:t")
        t.text = text
        t.set(qn("xml:space"), "preserve")
        r.append(t)
        h.append(r)
        p._p.append(h)

    # blocks
    def cover(self):
        d, L = self.doc, self.L
        def para(text, size, bold=False, before=0, after=0, color=None):
            p = d.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            f = p.paragraph_format
            f.space_before, f.space_after = Pt(before), Pt(after)
            f.line_spacing_rule = WD_LINE_SPACING.SINGLE
            r = p.add_run(text)
            r.font.size = Pt(size)
            r.font.bold = bold
            set_fonts(r._r.get_or_add_rPr(), L["latin_head"], L["head"])
            if color:
                r.font.color.rgb = color
            return p
        para(L["title"], 28, True, before=40, after=8, color=HEAD_COLOR)
        para(L["sub"], 14, after=24)
        art = os.path.join(ROOT, "docs", "media", "fdtp", "manual", COVER_ART)
        photo = art if os.path.exists(art) else os.path.join(ROOT, "docs", "media", "fdtp", "photos", COVER_PHOTO)
        p = d.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.add_run().add_picture(photo, width=Mm(110))
        # the boast, its note, then the "while I was at it" lines
        for i, line in enumerate(L["blurb"]):
            bright = [RGBColor(0xF0, 0x8A, 0x24), RGBColor(0x1E, 0x9B, 0xD7),
                      RGBColor(0x2B, 0xAE, 0x66), RGBColor(0x8E, 0x5B, 0xD9)]
            size, bold, color = ((11.5, True, RGBColor(0xE4, 0x40, 0x7B)) if i == 0 else
                                 (9.5, False, RGBColor(0x6B, 0x6B, 0x6B)) if i == 1 else
                                 (11, True, bright[(i - 2) % len(bright)]))
            para(line, size, bold=bold, before=(14 if i == 0 else 8 if i == 2 else 0), after=2, color=color)
        para(L["ver"].format(v=self.ver), 13, before=24, after=4)
        para(L["author"], 12).add_run().add_break(WD_BREAK.PAGE)

    def toc(self):
        d, L = self.doc, self.L
        p = d.add_paragraph()
        r = p.add_run(L["toc"])
        r.font.size = Pt(16)
        r.font.bold = True
        r.font.color.rgb = HEAD_COLOR
        set_fonts(r._r.get_or_add_rPr(), L["latin_head"], L["head"])
        p.paragraph_format.space_after = Pt(12)
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p = d.add_paragraph()
        self.field(p, 'TOC \\o "1-3" \\h \\z \\u', L["toc_hint"])

    def heading(self, text, level):
        self.doc.add_heading(text, level)

    def paragraph(self, text):
        p = self.doc.add_paragraph()
        self.inline(p, text)
        return p

    def bullet(self, text, level=0, number=None):
        p = self.doc.add_paragraph()
        f = p.paragraph_format
        indent = 6 + 6 * level
        f.left_indent = Mm(indent)
        f.first_line_indent = Mm(-5)
        f.space_after = Pt(3)
        mark = f"{number}." if number else "・" if self.lang == "ja" else "•"
        p.add_run(mark + "\t")
        tabs = f.tab_stops
        tabs.add_tab_stop(Mm(indent))
        self.inline(p, text)

    def code(self, lines):
        L = self.L
        for i, line in enumerate(lines):
            p = self.doc.add_paragraph()
            f = p.paragraph_format
            f.left_indent = Mm(4)
            f.space_before = Pt(4 if i == 0 else 0)
            f.space_after = Pt(6 if i == len(lines) - 1 else 0)
            f.line_spacing_rule = WD_LINE_SPACING.SINGLE
            f.keep_with_next = i < len(lines) - 1
            f.keep_together = True
            shade(p._p.get_or_add_pPr(), CODE_FILL)
            r = p.add_run(line if line else " ")
            r.font.size = Pt(9)
            set_fonts(r._r.get_or_add_rPr(), L["mono"], L["mono_ea"])

    def image(self, path, alt):
        if os.path.basename(path) == COVER_PHOTO:
            return
        want = PHOTO_W if path.lower().endswith(".jpg") else SHOT_W
        p = self.doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        f = p.paragraph_format
        f.keep_with_next = True
        f.space_before, f.space_after = Pt(6), Pt(2)
        f.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.add_run().add_picture(path, width=image_size(path, want))
        self.caption(alt)

    def caption(self, alt, cell=None):
        self.fig += 1
        text = self.L["fig"].format(n=self.fig, t=alt)
        p = cell.add_paragraph(style="Caption") if cell else self.doc.add_paragraph(style="Caption")
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
        p.add_run(text)

    def table(self, rows):
        L = self.L
        ncol = max(len(r) for r in rows)
        rows = [r + [""] * (ncol - len(r)) for r in rows]
        has_image = any(IMG.fullmatch(c.strip()) for r in rows for c in r)
        # column widths from the text widths (CJK counts 2), at least 18 mm
        def width(s):
            s = re.sub(r"[`*]", "", s)
            return sum(2 if ord(ch) > 0x2E7F else 1 for ch in s)
        if has_image:
            weights = [1] * ncol
        else:
            weights = []
            for c in range(ncol):
                ws = sorted(width(r[c]) for r in rows)
                weights.append(max(4, ws[-1] * 0.6 + ws[len(ws) // 2] * 0.4))
        total = sum(weights)
        mm = [max(18.0, TEXT_W * w / total) for w in weights]
        # Short columns (labels such as a machine's name) on one line:
        # about 1.75 mm per half-width unit at 9.5 pt, plus the padding.
        if not has_image:
            for c in range(ncol):
                longest = max(width(r[c]) for r in rows)
                if longest <= 24:
                    mm[c] = max(mm[c], longest * 1.75 + 3)
        scale = TEXT_W / sum(mm)
        mm = [m * scale for m in mm]

        t = self.doc.add_table(rows=len(rows), cols=ncol)
        t.style = self.doc.styles["Table Grid"]
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        t.autofit = False                            # w:tblLayout fixed
        tblpr = t._tbl.tblPr
        tblw = tblpr.find(qn("w:tblW"))
        if tblw is None:
            tblw = OxmlElement("w:tblW")
            tblpr.append(tblw)
        tblw.set(qn("w:type"), "dxa")
        tblw.set(qn("w:w"), str(int(sum(mm) * 56.6929)))
        for c, w in enumerate(mm):
            t.columns[c].width = Mm(w)
        header = not has_image
        for i, r in enumerate(rows):
            tr = t.rows[i]
            if i == 0 and header:
                trpr = tr._tr.get_or_add_trPr()
                th = OxmlElement("w:tblHeader")
                trpr.append(th)
            cant = OxmlElement("w:cantSplit")
            tr._tr.get_or_add_trPr().append(cant)
            for c, text in enumerate(r):
                cell = tr.cells[c]
                cell.width = Mm(mm[c])
                p = cell.paragraphs[0]
                f = p.paragraph_format
                f.space_after = Pt(0)
                if L["cell_exact"]:
                    f.line_spacing_rule = WD_LINE_SPACING.EXACTLY
                    f.line_spacing = Pt(L["cell_exact"])
                else:
                    f.line_spacing = 1.05
                m = IMG.fullmatch(text.strip())
                if m:
                    f.line_spacing_rule = WD_LINE_SPACING.SINGLE
                    f.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    path = self.img_path(m.group(2))
                    p.add_run().add_picture(path, width=image_size(path, mm[c] - 4))
                    self.caption(m.group(1), cell)
                    continue
                if i == 0 and header:
                    shade(cell._tc.get_or_add_tcPr(), HEADER_FILL)
                    self.inline(p, text, L["cell"], bold=True)
                else:
                    self.inline(p, text, L["cell"])
        # a little space after the table
        sp = self.doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(2)
        sp.paragraph_format.line_spacing_rule = WD_LINE_SPACING.EXACTLY
        sp.paragraph_format.line_spacing = Pt(6)

    def img_path(self, rel):
        path = os.path.normpath(os.path.join(SRC, rel))
        if not os.path.isfile(path):
            sys.exit(f"make_manual: image missing: {rel}")
        return path

    # the Markdown
    def build(self, md):
        lines = md.split("\n")
        i = 0
        while i < len(lines) and not lines[i].startswith("## "):
            i += 1
        self.cover()
        self.toc()
        para = []

        def flush():
            if para:
                self.paragraph(" ".join(para) if self.lang == "en" else "".join(para))
                para.clear()
        while i < len(lines):
            line = lines[i]
            s = line.strip()
            if s.startswith("```"):
                flush()
                block = []
                i += 1
                while i < len(lines) and not lines[i].strip().startswith("```"):
                    block.append(lines[i])
                    i += 1
                self.code(block)
            elif s == "---" or s == "":
                flush()
            elif s.startswith("#"):
                flush()
                level = len(s) - len(s.lstrip("#"))
                self.heading(s[level:].strip(), level - 1)
            elif s.startswith("|"):
                flush()
                rows = []
                while i < len(lines) and lines[i].strip().startswith("|"):
                    cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                    if not all(re.fullmatch(r":?-{3,}:?", c) for c in cells):
                        rows.append(cells)
                    i += 1
                self.table(rows)
                continue
            elif IMG.fullmatch(s):
                flush()
                m = IMG.fullmatch(s)
                self.image(self.img_path(m.group(2)), m.group(1))
            elif re.match(r"- ", s):
                flush()
                self.bullet(s[2:])
            elif re.match(r"\d+\. ", s):
                flush()
                n, text = s.split(". ", 1)
                self.bullet(text, 1 if line.startswith("  ") else 0, number=n)
            else:
                para.append(s)
            i += 1
        flush()


IMG = re.compile(r"!\[([^\]]*)\]\(([^)]+)\)")


def main(argv):
    langs = argv or ["ja", "en"]
    ver = version()
    for lang in langs:
        if lang not in LANG:
            sys.exit(f"make_manual: unknown language {lang}")
        md = open(os.path.join(SRC, f"FDTP_manual_{lang}.md"), encoding="utf-8").read()
        b = Builder(lang, ver)
        b.build(md)
        out_dir = os.path.join(OUT, lang)
        os.makedirs(out_dir, exist_ok=True)
        out = os.path.join(out_dir, f"FDTP_manual_{lang}.docx")
        b.doc.save(out)
        print(f"make_manual: {os.path.relpath(out, ROOT)} ({b.fig} figures, version {ver})")


if __name__ == "__main__":
    main(sys.argv[1:])
